#!/usr/bin/env python3
"""Insert embedded-video slides into an existing lecture deck.

Each new slide is cloned from a figure slide of the SAME deck (title, body
text boxes, caption, footer citation and footer number keep the deck's own
positions, fonts, sizes and colors). The figure is replaced by an embedded
MP4 (with a poster frame), and speaker notes are added.

Original slides are never edited: the output package copies every original
part byte-for-byte from the input and only adds new parts plus the updated
slide order (ppt/presentation.xml, its rels, and [Content_Types].xml).

Usage:
    python insert_video_slides.py IN.pptx SPEC.json OUT.pptx

SPEC.json:
{
  "template_slide": 12,            # optional: original slide (1-based) to clone
  "inserts": [
    {"after": 17, "label": "17a", "title": "...",
     "paras": ["... **bold** ... *italic* ...", "..."],
     "caption": "...", "cite": "...", "notes": ["...", "..."],
     "video": "videos/01_x.mp4", "poster_t": 5.0,
     "template_slide": 17}          # optional per-insert override
  ]
}
"""
import copy
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A = "{%s}" % NS["a"]
R_ID = "{%s}id" % NS["r"]
EMU_IN = 914400


# ---------------------------------------------------------------- analysis
def classify(slide, slide_w, slide_h):
    """Return roles for the shapes of a content slide."""
    texts, pics, other = [], [], []
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pics.append(sh)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh)
        else:
            other.append(sh)
    footer_top = slide_h * 0.9
    footer = [s for s in texts if s.top >= footer_top]
    number = [s for s in footer if re.fullmatch(r"\d{1,3}[a-z]?", s.text_frame.text.strip())]
    cite = [s for s in footer if s not in number]
    upper = [s for s in texts if s.top < footer_top]
    title = min(upper, key=lambda s: (s.top, -s.width)) if upper else None
    big_pic = max(pics, key=lambda p: p.width * p.height) if pics else None
    captions = []
    for p in pics:
        below = [s for s in upper if s is not title
                 and p.top + p.height - EMU_IN * 0.1 <= s.top <= p.top + p.height + EMU_IN * 0.9
                 and s.left < p.left + p.width and s.left + s.width > p.left]
        if below:
            captions.append((p, min(below, key=lambda s: s.top)))
    cap_shapes = [c for _, c in captions]
    big_caption = next((c for p, c in captions if p is big_pic), None)
    body = [s for s in upper if s is not title and s not in cap_shapes]
    body.sort(key=lambda s: (s.top, s.left))
    return dict(title=title, body=body, pics=pics, big_pic=big_pic,
                caption=big_caption, captions=cap_shapes, cite=cite[0] if cite else None,
                number=number[0] if number else None, other=other)


def pick_template(prs, slide_w, slide_h):
    """First content slide with exactly one picture that has a caption and body text."""
    n = len(prs.slides)
    best = None
    for i in range(1, n - 1):
        roles = classify(prs.slides[i], slide_w, slide_h)
        if roles["big_pic"] is None or roles["caption"] is None or not roles["body"]:
            continue
        if roles["title"] is None or roles["cite"] is None or roles["number"] is None:
            continue
        if len(roles["pics"]) == 1:
            return i
        best = best if best is not None else i
    if best is None:
        raise SystemExit("No content slide with figure + caption + footer found to clone")
    return best


# ---------------------------------------------------------------- text helpers
def _runs_from_markup(text):
    """Split '**bold**' / '*italic*' markup into (text, bold, italic) runs."""
    out = []
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            out.append((tok[2:-2], True, False))
        elif tok.startswith("*"):
            out.append((tok[1:-1], False, True))
        else:
            out.append((tok, False, False))
    return out


def set_paragraph(p_el, text, base_rpr):
    for r in p_el.findall(A + "r") + p_el.findall(A + "br") + p_el.findall(A + "fld"):
        p_el.remove(r)
    end = p_el.find(A + "endParaRPr")
    for t, bold, ital in _runs_from_markup(text):
        r = etree.SubElement(p_el, A + "r")
        rpr = copy.deepcopy(base_rpr)
        if bold:
            rpr.set("b", "1")
        if ital:
            rpr.set("i", "1")
        for h in rpr.findall(A + "hlinkClick"):
            rpr.remove(h)
        r.append(rpr)
        etree.SubElement(r, A + "t").text = t
        if end is not None:
            end.addprevious(r)


def plain_rpr(sp_el):
    """rPr of the first non-bold run (falls back to the first run)."""
    rprs = sp_el.findall(".//" + A + "r/" + A + "rPr")
    for rpr in rprs:
        if rpr.get("b") in (None, "0"):
            return rpr
    return rprs[0]


def set_shape_text(sp_el, paragraphs, body=False):
    """Replace a text shape's paragraphs, reusing its first paragraph's formatting.

    Body text starts from the deck's plain (non-bold) run; titles, captions and
    footers keep their first run's formatting (e.g. bold titles)."""
    tx = sp_el.find(".//{%s}txBody" % NS["p"])
    ps = tx.findall(A + "p")
    first = sp_el.find(".//" + A + "r/" + A + "rPr")
    base_p, base_rpr = ps[0], copy.deepcopy(plain_rpr(sp_el) if body else first)
    for p in ps:
        tx.remove(p)
    for text in paragraphs:
        p = copy.deepcopy(base_p)
        set_paragraph(p, text, base_rpr)
        tx.append(p)


def xfrm(sp_el):
    off = sp_el.find(".//" + A + "xfrm/" + A + "off")
    ext = sp_el.find(".//" + A + "xfrm/" + A + "ext")
    return off, ext


# ---------------------------------------------------------------- build
def poster_frame(video, t, out_png):
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", str(t), "-i", video,
                    "-frames:v", "1", out_png], check=True)


def video_size(video):
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0",
                          "-show_entries", "stream=width,height", "-of", "csv=p=0", video],
                         check=True, capture_output=True, text=True).stdout.strip()
    w, h = out.split(",")[:2]
    return int(w), int(h)


def build_slide(prs, tmpl_idx, ins, tmpdir, slide_w, slide_h):
    src = prs.slides[tmpl_idx]
    roles = classify(src, slide_w, slide_h)
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree)[2:]:            # drop layout placeholders
        tree.remove(el)
    bg = src._element.find("{%s}cSld/{%s}bg" % (NS["p"], NS["p"]))
    if bg is not None:
        new._element.find("{%s}cSld" % NS["p"]).insert(0, copy.deepcopy(bg))

    # Decorative non-text, non-picture shapes (rule lines etc.) are copied as-is.
    for sh in roles["other"]:
        if sh._element.find(".//" + A + "blip") is None:
            tree.append(copy.deepcopy(sh._element))

    def clone(sh):
        el = copy.deepcopy(sh._element)
        for node in el.iter():           # strip relationship-bearing attributes
            if R_ID in node.attrib:
                del node.attrib[R_ID]
        for h in el.findall(".//" + A + "hlinkClick"):
            h.getparent().remove(h)
        tree.append(el)
        return el

    set_shape_text(clone(roles["title"]), [ins["title"]])

    body = roles["body"]
    paras = ins["paras"]
    if len(body) == 1:
        set_shape_text(clone(body[0]), paras, body=True)
    else:
        use = body[:len(paras)]
        for i, sh in enumerate(use):
            chunk = paras[i:i + 1] if i < len(use) - 1 else paras[i:]
            set_shape_text(clone(sh), chunk, body=True)

    cap_el = clone(roles["caption"])
    set_shape_text(cap_el, [ins["caption"]])
    set_shape_text(clone(roles["cite"]), [ins["cite"]])
    num_el = clone(roles["number"])
    set_shape_text(num_el, [ins["label"]])
    # Keep the footer number right-aligned to the same right edge; widen for the letter.
    off, ext = xfrm(num_el)
    right = int(off.get("x")) + int(ext.get("cx"))
    cx = max(int(ext.get("cx")), int(int(ext.get("cx")) * len(ins["label"]) / max(1, len(roles["number"].text_frame.text.strip())) * 1.15))
    ext.set("cx", str(cx))
    off.set("x", str(right - cx))
    bp = num_el.find(".//" + A + "bodyPr")
    bp.set("wrap", "none")
    for p in num_el.findall(".//" + A + "pPr"):
        p.set("algn", "r")

    # Video in the figure's frame: fit preserving aspect ratio, top-aligned, centered.
    pic = roles["big_pic"]
    vw, vh = video_size(ins["video"])
    box_w, box_h = pic.width, pic.height
    cap_off, _ = xfrm(cap_el)
    box_h = min(box_h, int(cap_off.get("y")) - pic.top - int(EMU_IN * 0.08))
    scale = min(box_w / vw, box_h / vh)
    w, h = int(vw * scale), int(vh * scale)
    left = pic.left + (box_w - w) // 2
    top = pic.top
    png = os.path.join(tmpdir, os.path.basename(ins["video"]) + ".png")
    poster_frame(ins["video"], ins.get("poster_t", 3), png)
    mov = new.shapes.add_movie(ins["video"], left, top, w, h, poster_frame_image=png,
                               mime_type="video/mp4")
    mov.name = "Video: " + os.path.basename(ins["video"])
    # Caption sits directly under the video, same left edge as the video.
    cap_off.set("y", str(top + h + int(EMU_IN * 0.1)))
    cap_off.set("x", str(left))
    _, cap_ext = xfrm(cap_el)
    cap_ext.set("cx", str(max(w, int(cap_ext.get("cx")))))

    # Speaker notes, written the way the deck writes them ("• " bullets).
    tf = new.notes_slide.notes_text_frame
    tf.text = "\n\n".join("• " + n for n in ins["notes"]) + (
        "\n\nReferences:\n\n" + ins["notes_refs"] if ins.get("notes_refs") else "")
    return new


def reorder(prs, orig_ids, plan):
    """plan: list of (after_original_number, new_sldId_element) in insertion order."""
    lst = prs.slides._sldIdLst
    for after, el in plan:
        lst.remove(el)
    anchor = {}
    for after, el in plan:
        prev = anchor.get(after, orig_ids[after - 1])
        prev.addnext(el)
        anchor[after] = el


def rebuild_zip(orig_path, tmp_path, out_path):
    keep_new = {"ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"}
    with zipfile.ZipFile(orig_path) as zo, zipfile.ZipFile(tmp_path) as zn, \
            zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zw:
        orig_names = zo.namelist()
        for name in orig_names:
            data = zn.read(name) if name in keep_new else zo.read(name)
            zw.writestr(zo.getinfo(name) if name not in keep_new else name, data)
        for name in zn.namelist():
            if name not in orig_names:
                ctype = zipfile.ZIP_STORED if name.endswith(".mp4") else zipfile.ZIP_DEFLATED
                zw.writestr(zipfile.ZipInfo(name), zn.read(name), compress_type=ctype)


def main(in_path, spec_path, out_path):
    spec = json.load(open(spec_path))
    base = os.path.dirname(os.path.abspath(spec_path))
    prs = Presentation(in_path)
    slide_w, slide_h = prs.slide_width, prs.slide_height
    orig_ids = list(prs.slides._sldIdLst)
    n_orig = len(orig_ids)
    default_t = spec.get("template_slide")
    default_t = default_t - 1 if default_t else pick_template(prs, slide_w, slide_h)
    plan = []
    with tempfile.TemporaryDirectory() as td:
        for ins in spec["inserts"]:
            if not 1 <= ins["after"] < n_orig:
                raise SystemExit(f"'after' must be between 1 and {n_orig - 1}: {ins['after']}")
            ins = dict(ins, video=os.path.join(base, ins["video"]))
            t = ins.get("template_slide")
            build_slide(prs, (t - 1) if t else default_t, ins, td, slide_w, slide_h)
            plan.append((ins["after"], prs.slides._sldIdLst[-1]))
        reorder(prs, orig_ids, plan)
        tmp = os.path.join(td, "tmp.pptx")
        prs.save(tmp)
        rebuild_zip(in_path, tmp, out_path)
    print(f"{out_path}: {n_orig} original + {len(plan)} video slides = {n_orig + len(plan)}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
